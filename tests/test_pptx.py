"""MS-910: text extraction and picture OCR for PowerPoint decks.

Decks are built with python-pptx inside each test, so no fixture files and no
database are needed."""

import io
import os
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

from config import config
from services.ingest import documents
from services.ingest.documents import DOCUMENT_EXTRACTORS, _extract_pdf_with_progress, extract_text_from_pptx
from services.ingest.ocr import SlideOcr, describe_ocr_outcome, new_ocr_stats

LONG = "Kebijakan perjalanan dinas berlaku untuk seluruh karyawan tetap dan kontrak. " * 4


class PptxExtractionTests(unittest.TestCase):
    def setUp(self):
        self._dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._dir.cleanup)
        self.deck = Presentation()
        self.blank = self.deck.slide_layouts[6]

    def add_slide(self, text=None):
        slide = self.deck.slides.add_slide(self.blank)
        if text:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = text
        return slide

    def extract(self):
        path = os.path.join(self._dir.name, "deck.pptx")
        self.deck.save(path)
        return extract_text_from_pptx(path)

    def test_text_boxes_tables_groups_and_notes_keep_their_slide_number(self):
        self.add_slide(LONG)
        table_slide = self.add_slide(LONG)
        table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "Vendor", "Nilai kontrak"
        table.cell(1, 0).text, table.cell(1, 1).text = "PT Sinar Abadi", "Rp 120 juta"
        group_slide = self.add_slide(LONG)
        group_slide.shapes.add_group_shape().shapes.add_textbox(
            Inches(1), Inches(3), Inches(4), Inches(1)).text_frame.text = "SLA 99,5 persen"
        group_slide.notes_slide.notes_text_frame.text = "Kode proyek ZETA-42"

        pages = self.extract()

        self.assertEqual([p["page"] for p in pages], [1, 2, 3])
        self.assertIn("PT Sinar Abadi | Rp 120 juta", pages[1]["text"])
        self.assertIn("SLA 99,5 persen", pages[2]["text"])
        self.assertIn("Kode proyek ZETA-42", pages[2]["text"])

    def test_hidden_slides_are_left_out(self):
        self.add_slide(LONG)
        self.add_slide("RAHASIA " + LONG)._element.set("show", "0")
        self.add_slide(LONG)

        pages = self.extract()

        self.assertEqual([p["page"] for p in pages], [1, 3])
        self.assertNotIn("RAHASIA", " ".join(p["text"] for p in pages))

    def test_title_slides_are_carried_into_the_next_slide_and_cited_as_it(self):
        self.add_slide("Agenda")
        self.add_slide("Bagian 1")
        self.add_slide(LONG)
        self.add_slide(LONG)

        pages = self.extract()

        self.assertEqual([p["page"] for p in pages], [3, 4])
        self.assertTrue(pages[0]["text"].startswith("Agenda Bagian 1 Kebijakan"))

    def test_a_short_last_slide_joins_the_one_before(self):
        self.add_slide(LONG)
        self.add_slide("Terima kasih")

        pages = self.extract()

        self.assertEqual([p["page"] for p in pages], [1])
        self.assertTrue(pages[0]["text"].endswith("Terima kasih"))

    def test_a_slide_with_one_sentence_of_content_keeps_its_own_number(self):
        self.add_slide(LONG)
        self.add_slide("Batas anggaran proyek Rp 75 juta disetujui direksi pada rapat Maret 2026 di Jakarta.")
        self.add_slide(LONG)

        pages = self.extract()

        self.assertEqual([p["page"] for p in pages], [1, 2, 3])

    def test_a_broken_slide_is_skipped_and_the_rest_kept(self):
        for _ in range(3):
            self.add_slide(LONG)
        real = documents._pptx_slide_text
        calls = []

        def fail_on_second(slide, *args):
            calls.append(slide)
            if len(calls) == 2:
                raise ValueError("bad shape")
            return real(slide, *args)

        with patch.object(documents, "_pptx_slide_text", fail_on_second):
            pages = self.extract()

        self.assertEqual([p["page"] for p in pages], [1, 3])

    def test_deck_without_text_and_non_pptx_files_give_none(self):
        self.add_slide()
        self.assertIsNone(self.extract())
        legacy = os.path.join(self._dir.name, "old.pptx")
        with open(legacy, "wb") as handle:
            handle.write(b"\xd0\xcf\x11\xe0 not a zip package")
        self.assertIsNone(extract_text_from_pptx(legacy))

    def test_pptx_is_an_accepted_upload_type(self):
        self.assertIs(DOCUMENT_EXTRACTORS[".pptx"], extract_text_from_pptx)
        self.assertNotIn(".ppt", DOCUMENT_EXTRACTORS)


def _png(width, height):
    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (255, 255, 255)).save(buffer, "PNG")
    buffer.seek(0)
    return buffer


class PptxPictureOcrTests(unittest.TestCase):
    """The OCR engine is faked (as in test_ocr.py), so these run without
    tesseract; the real engine on slide pictures is covered by the golden set."""

    def setUp(self):
        self._dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._dir.cleanup)
        self.path = os.path.join(self._dir.name, "deck.pptx")
        self.deck = Presentation()
        for name, value in {
            "ocr_enabled": True, "ocr_max_pages_per_file": 40,
            "ocr_time_limit_seconds": 120, "ocr_min_picture_pixels": 128 * 128,
        }.items():
            patcher = patch.object(config, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def add_picture_slide(self, width=800, height=300, grouped=False, text=LONG):
        slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(6), Inches(1)).text_frame.text = text
        shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
        shapes.add_picture(_png(width, height), Inches(1), Inches(2), width=Inches(4))

    def extract(self, ocr_text="Batas anggaran Rp 75 juta", **kwargs):
        self.deck.save(self.path)
        stats = new_ocr_stats()
        with patch.object(SlideOcr, "_ocr", return_value=ocr_text) as fake:
            pages = extract_text_from_pptx(self.path, ocr_stats=stats, **kwargs)
        return pages, stats, fake

    def test_text_in_pictures_is_indexed_under_its_slide(self):
        self.add_picture_slide()
        self.add_picture_slide(grouped=True)
        pages, stats, fake = self.extract()
        self.assertEqual(fake.call_count, 2)
        self.assertEqual([p["page"] for p in pages], [1, 2])
        self.assertTrue(all("Batas anggaran Rp 75 juta" in p["text"] for p in pages))
        self.assertEqual(stats["read"], 2)

    def test_icons_and_logos_are_not_ocrd(self):
        self.add_picture_slide(width=64, height=64)
        pages, _, fake = self.extract()
        fake.assert_not_called()
        self.assertEqual(len(pages), 1)

    def test_pictures_without_text_are_not_reported_as_unreadable(self):
        self.add_picture_slide()
        pages, stats, _ = self.extract(ocr_text="")
        self.assertEqual(stats["empty"], 0)
        self.assertIsNone(describe_ocr_outcome(stats))
        self.assertEqual(len(pages), 1)

    def test_picture_cap_is_shared_with_pdf_and_reported(self):
        for _ in range(3):
            self.add_picture_slide()
        with patch.object(config, "ocr_max_pages_per_file", 2):
            _, stats, fake = self.extract()
        self.assertEqual(fake.call_count, 2)
        self.assertEqual(stats["over_limit"], 1)

    def test_ocr_can_be_switched_off(self):
        self.add_picture_slide()
        with patch.object(config, "ocr_enabled", False):
            _, _, fake = self.extract()
        fake.assert_not_called()

    def test_slides_being_ocrd_report_the_ocr_stage(self):
        self.add_picture_slide()
        self.add_picture_slide(width=64, height=64)
        self.deck.save(self.path)
        events = []
        with patch.object(SlideOcr, "_ocr", return_value="Batas anggaran Rp 75 juta"):
            _extract_pdf_with_progress(
                extract_text_from_pptx, self.path, 0, 1,
                lambda *args: events.append(args), new_ocr_stats(),
            )
        stages = [(e[0], e[2] if len(e) > 2 else None) for e in events]
        self.assertIn(("ocr", {"page": 1, "total": 2}), stages)
        self.assertNotIn(("ocr", {"page": 2, "total": 2}), stages)
        progress = [e[1] for e in events]
        self.assertEqual(progress, sorted(progress))


if __name__ == "__main__":
    unittest.main()
